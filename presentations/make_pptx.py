#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Generate two PPTX decks for the QR sorting robot:
   1) Presentation-Simple.pptx     — for a general audience
   2) Presentation-Technical.pptx  — detailed (instructor / evaluation)

   Pipeline described: PC-decode. The ESP32-CAM streams MJPEG over WiFi, a PC
   (Python + OpenCV/pyzbar) decodes the QR and forwards "QR:<code>" to the brain
   over USB serial.
"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

# ---- color theme ----
NAVY   = RGBColor(0x10, 0x2A, 0x43)   # title / divider background
BLUE   = RGBColor(0x1F, 0x6F, 0xB2)   # accent
LBLUE  = RGBColor(0xE8, 0xF1, 0xF9)   # box background
DARK   = RGBColor(0x22, 0x2B, 0x33)   # main text
GRAY   = RGBColor(0x5A, 0x66, 0x70)   # secondary text
WHITE  = RGBColor(0xFF, 0xFF, 0xFF)
GREEN  = RGBColor(0x2E, 0x8B, 0x57)
AMBER  = RGBColor(0xC8, 0x7A, 0x00)

EMU_W, EMU_H = Inches(13.333), Inches(7.5)

def new_deck():
    prs = Presentation()
    prs.slide_width  = EMU_W
    prs.slide_height = EMU_H
    return prs

def blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])

def rect(slide, x, y, w, h, color, line=None):
    from pptx.enum.shapes import MSO_SHAPE
    sp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    sp.fill.solid(); sp.fill.fore_color.rgb = color
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line; sp.line.width = Pt(1)
    sp.shadow.inherit = False
    return sp

def txt(slide, x, y, w, h, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
        space_after=6, line_spacing=1.05):
    """runs: list of paragraphs; each paragraph = list of (text, size, color, bold)."""
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame; tf.word_wrap = True
    tf.vertical_anchor = anchor
    for i, para in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align; p.space_after = Pt(space_after); p.line_spacing = line_spacing
        for (t, sz, col, bold) in para:
            r = p.add_run(); r.text = t
            r.font.size = Pt(sz); r.font.color.rgb = col; r.font.bold = bold
            r.font.name = "Calibri"
    return tb

def bullets(slide, x, y, w, h, items, size=20, color=DARK, gap=10):
    """items: list of (text, level) or (text, level, color)."""
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame; tf.word_wrap = True
    for i, it in enumerate(items):
        t, lvl = it[0], it[1]
        col = it[2] if len(it) > 2 else color
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(gap); p.line_spacing = 1.05
        bullet = "•  " if lvl == 0 else "–  "
        r = p.add_run(); r.text = bullet + t
        r.font.size = Pt(size if lvl == 0 else size-2)
        r.font.color.rgb = col
        r.font.bold = (lvl == 0 and len(it) > 2)
        r.font.name = "Calibri"
        p.level = lvl
        if lvl > 0:
            pPr = p._pPr if p._pPr is not None else p.get_or_add_pPr()
            pPr.set("marL", str(Inches(0.5).emu)); pPr.set("indent", "0")
    return tb

# ---------- slide types ----------
def slide_title(prs, kicker, title, subtitle, footer):
    s = blank(prs)
    rect(s, 0, 0, EMU_W, EMU_H, NAVY)
    rect(s, 0, Inches(3.55), EMU_W, Inches(0.06), BLUE)
    txt(s, Inches(0.9), Inches(1.25), Inches(11.5), Inches(0.6),
        [[(kicker, 18, BLUE, True)]])
    txt(s, Inches(0.9), Inches(1.75), Inches(11.5), Inches(1.7),
        [[(title, 40, WHITE, True)]])
    txt(s, Inches(0.9), Inches(3.75), Inches(11.5), Inches(1.6),
        [[(subtitle, 20, RGBColor(0xC4,0xD3,0xE0), False)]])
    txt(s, Inches(0.9), Inches(6.7), Inches(11.5), Inches(0.5),
        [[(footer, 13, RGBColor(0x90,0xA4,0xB5), False)]])
    return s

def slide_section(prs, num, title):
    s = blank(prs)
    rect(s, 0, 0, EMU_W, EMU_H, NAVY)
    rect(s, Inches(0.9), Inches(2.7), Inches(1.4), Inches(1.4), BLUE)
    txt(s, Inches(0.9), Inches(2.7), Inches(1.4), Inches(1.4),
        [[(num, 44, WHITE, True)]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    txt(s, Inches(2.6), Inches(2.7), Inches(10.0), Inches(1.4),
        [[(title, 34, WHITE, True)]], anchor=MSO_ANCHOR.MIDDLE)
    return s

def header(s, title, sub=None):
    rect(s, 0, 0, EMU_W, Inches(1.15), NAVY)
    rect(s, 0, Inches(1.15), EMU_W, Inches(0.06), BLUE)
    txt(s, Inches(0.7), Inches(0.18), Inches(12), Inches(0.85),
        [[(title, 28, WHITE, True)]], anchor=MSO_ANCHOR.MIDDLE)
    if sub:
        txt(s, Inches(0.72), Inches(1.28), Inches(12), Inches(0.4),
            [[(sub, 15, GRAY, False)]])

def slide_bullets(prs, title, items, sub=None, size=20):
    s = blank(prs); header(s, title, sub)
    bullets(s, Inches(0.8), Inches(1.7), Inches(11.7), Inches(5.4), items, size=size)
    return s

def callout(slide, x, y, w, h, head, body, color=LBLUE, headcol=BLUE, bodycol=DARK):
    rect(slide, x, y, w, h, color)
    rect(slide, x, y, Inches(0.12), h, headcol)
    txt(slide, x+Inches(0.35), y+Inches(0.2), w-Inches(0.6), h-Inches(0.4),
        [[(head, 18, headcol, True)], [(body, 16, bodycol, False)]], space_after=6)

def add_table(slide, x, y, w, rows, col_w=None, head_fill=BLUE, fsize=14):
    nrows = len(rows); ncols = len(rows[0])
    h = Inches(0.45) * nrows
    gtbl = slide.shapes.add_table(nrows, ncols, x, y, w, h).table
    if col_w:
        for j, cw in enumerate(col_w):
            gtbl.columns[j].width = cw
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            c = gtbl.cell(i, j)
            c.margin_left = Inches(0.08); c.margin_right = Inches(0.08)
            c.margin_top = Inches(0.02); c.margin_bottom = Inches(0.02)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = c.text_frame; tf.word_wrap = True
            p = tf.paragraphs[0]; r = p.add_run(); r.text = str(val)
            r.font.name = "Calibri"
            if i == 0:
                c.fill.solid(); c.fill.fore_color.rgb = head_fill
                r.font.size = Pt(fsize+1); r.font.bold = True; r.font.color.rgb = WHITE
            else:
                c.fill.solid(); c.fill.fore_color.rgb = WHITE if i % 2 else LBLUE
                r.font.size = Pt(fsize); r.font.color.rgb = DARK
    return gtbl

def save(prs, name):
    out = os.path.join(OUTDIR, name)
    prs.save(out)
    print("written:", out, "(", len(prs.slides._sldIdLst), "slides )")

OUTDIR = os.path.expanduser("~/repos/Sorting-robot/presentations")
os.makedirs(OUTDIR, exist_ok=True)

FOOT = "QR Sorting Robot  •  Cartesian XYZ gantry + ESP32  •  PC-decode pipeline"

# =====================================================================
#  DECK 1 — SIMPLE (general audience)
# =====================================================================
def deck_simple():
    p = new_deck()

    slide_title(p, "ROBOTICS PROJECT",
                "The robot that sorts boxes by their QR code",
                "How it picks up a box, reads its code and places it on the right shelf "
                "by itself — explained simply, step by step.",
                FOOT)

    slide_bullets(p, "What does the robot do?",
        [("Boxes arrive, each with a QR code stuck on it", 0),
         ("The code says WHERE the box must go (e.g. MS0011, CJ0012)", 0),
         ("The robot picks up the box and reads its code with a camera", 0),
         ("It decides by itself which shelf it belongs to and places it there", 0),
         ("It returns and waits for the next box — fully automatic", 0)],
        sub="A miniature automated warehouse: pick & place with sorting by code")

    # visual 5-step flow
    s = blank(p); header(s, "In short: 5 steps")
    steps = [("1", "PICK the box\nfrom the feed zone"),
             ("2", "BRING the box\nto the camera"),
             ("3", "READ\nthe QR code"),
             ("4", "DECIDE\nthe right shelf"),
             ("5", "PLACE the box\nand return")]
    bx = Inches(0.55); bw = Inches(2.35); gap = Inches(0.12)
    for i,(n,t) in enumerate(steps):
        x = bx + i*(bw+gap)
        rect(s, x, Inches(2.6), bw, Inches(2.2), LBLUE)
        rect(s, x, Inches(2.6), bw, Inches(0.65), BLUE)
        txt(s, x, Inches(2.6), bw, Inches(0.65), [[(n, 24, WHITE, True)]],
            align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        txt(s, x+Inches(0.1), Inches(3.4), bw-Inches(0.2), Inches(1.3),
            [[(t, 16, DARK, False)]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

    slide_bullets(p, "How the robot moves",
        [("It has 3 axes sliding on straight rails: X, Y and Z", 0),
         ("X = left/right, Y = front/back, Z = up/down", 1),
         ("Combining the 3, the tip reaches any point (like a 3D map)", 0),
         ("Same principle as a 3D printer or a CNC machine", 0),
         ("Stepper motors: they turn in equal steps → very precise positioning", 0)],
        sub="A Cartesian XYZ frame (gantry), not an articulated arm")

    slide_bullets(p, "How the robot “sees” the QR code",
        [("On its tip it carries a small camera (ESP32-CAM)", 0),
         ("The camera STREAMS the video over WiFi to a PC", 0),
         ("On the PC, a Python program (OpenCV) decodes the code into text", 0),
         ("Why on the PC? It is far more accurate than decoding on the tiny chip", 1),
         ("The camera is always brought to the same spot so it stays in focus", 0)],
        sub="A QR code = a “2D barcode” that hides a piece of text")

    # communication
    s = blank(p); header(s, "The three actors of the system")
    callout(s, Inches(0.5), Inches(1.7), Inches(4.0), Inches(2.1),
            "THE EYE — ESP32-CAM",
            "Sits on the tip. Streams the video over WiFi. Decodes nothing, moves nothing.")
    callout(s, Inches(4.7), Inches(1.7), Inches(4.0), Inches(2.1),
            "THE PC",
            "Receives the video, decodes the QR with OpenCV, sends the code to the brain.",
            color=RGBColor(0xFD,0xF3,0xE0), headcol=AMBER)
    callout(s, Inches(8.9), Inches(1.7), Inches(3.9), Inches(2.1),
            "THE BRAIN — ESP32",
            "Gets the code, drives the motors and the hand, decides the shelf.",
            color=RGBColor(0xEA,0xF7,0xEE), headcol=GREEN)
    txt(s, Inches(0.7), Inches(4.2), Inches(12), Inches(1.6),
        [[("Data flows one way:  ", 20, DARK, True),
          ("camera  →  (WiFi video)  →  PC  →  (USB cable, “QR:CJ0012”)  →  brain.",
           20, BLUE, True)]], line_spacing=1.1)

    slide_bullets(p, "How it decides where to put the box",
        [("The shelves are a grid of 3 rows × 2 columns", 0),
         ("Each code has a fixed cell, set in advance:", 0),
         ("MS0011 → row 1 left, CJ0012 → row 2 right, etc.", 1),
         ("The robot looks the code up in a table and goes exactly there", 0),
         ("Unknown code? It rejects it — it will not risk a mistake", 0)],
        sub="Fixed code → cell mapping: simple and safe")

    slide_bullets(p, "Why ESP32 and not classic Arduino?",
        [("It has WiFi/radio built in → needed for the wireless video", 0),
         ("It has far more memory → it can handle the camera image", 0),
         ("It has 2 cores → it moves the motors AND communicates at the same time", 0),
         ("It generates the motor pulses with dedicated hardware (precise)", 0),
         ("An Arduino Uno (2KB memory, no radio) simply could not", 0, AMBER)],
        sub="ESP32 = a small, cheap and powerful “mini-computer”")

    slide_bullets(p, "Safety: how it does NOT hit the shelves",
        [("It ALWAYS raises the arm (Z) before moving sideways", 0),
         ("On power-up it does “homing”: it finds the zero point by itself", 0),
         ("It has software limits — it cannot leave the working area", 0),
         ("A STOP command that halts everything instantly", 0)],
        sub="Simple rules that prevent collisions")

    slide_bullets(p, "In short — what we built",
        [("An XYZ robot that sorts boxes fully automatically, by QR code", 0, BLUE),
         ("Two ESP32 boards: one “sees”, one “drives”", 0),
         ("A PC decodes the code for high accuracy", 0),
         ("Precise motors + camera + a hand that grabs the box", 0),
         ("Every decision was made for reliability and simplicity", 0, GREEN)],
        sub="Thank you! Questions?")

    save(p, "Presentation-Simple.pptx")

# =====================================================================
#  DECK 2 — TECHNICAL (instructor / evaluation)
# =====================================================================
def deck_technical():
    p = new_deck()

    slide_title(p, "TECHNICAL DOCUMENTATION",
                "QR Sorting Robot — Cartesian XYZ gantry",
                "Architecture, motion control, vision, the PC-decode pipeline, "
                "the sorting logic and the design decisions.",
                FOOT)

    # AGENDA
    slide_bullets(p, "Table of contents",
        [("1.  System architecture (two MCUs + a PC)", 0),
         ("2.  Mechanics: XYZ gantry + screw actuators", 0),
         ("3.  Motion control: steppers, A4988, FastAccelStepper", 0),
         ("4.  Coordinates, homing, software limits", 0),
         ("5.  Gripper (SG90 / LEDC PWM)", 0),
         ("6.  Vision: PC-decode pipeline (ESP32-CAM stream + OpenCV)", 0),
         ("7.  Communication (WiFi stream + USB serial)", 0),
         ("8.  Sorting logic (CODE_MAP) + state machine", 0),
         ("9.  Power & pin map", 0),
         ("10. Design decisions & lessons", 0)], size=18)

    # 1. architecture
    slide_section(p, "1", "System architecture")
    s = blank(p); header(s, "Two microcontrollers + a PC, separate responsibilities")
    callout(s, Inches(0.5), Inches(1.4), Inches(4.0), Inches(2.6),
            "ESP32-CAM — VISION",
            "OV2640 + PSRAM. Streams MJPEG over WiFi (esp32cam_stream). Decodes nothing, "
            "controls nothing mechanical. Rides on the Z carriage.")
    callout(s, Inches(4.7), Inches(1.4), Inches(4.0), Inches(2.6),
            "PC — DECODER",
            "decode_qr.py: reads the MJPEG stream, decodes the QR with OpenCV/pyzbar, "
            "forwards “QR:<code>” over USB serial.",
            color=RGBColor(0xFD,0xF3,0xE0), headcol=AMBER)
    callout(s, Inches(8.9), Inches(1.4), Inches(3.9), Inches(2.6),
            "ESP32 WROOM-32D — BRAIN",
            "State machine (brainPC). STEP/DIR for 3× A4988, homing on endstops, "
            "servo, receives QR over serial, decides the cell.",
            color=RGBColor(0xEA,0xF7,0xEE), headcol=GREEN)
    callout(s, Inches(0.7), Inches(4.4), Inches(11.9), Inches(1.9),
            "Why the split?",
            "The camera takes up nearly all ESP32-CAM pins, and reliable QR decoding is heavy. "
            "The motion loop is timing-critical and must never be blocked. → camera streams, "
            "PC decodes (accurately), brain moves.",
            color=RGBColor(0xFD,0xF3,0xE0), headcol=AMBER)

    slide_bullets(p, "Why ESP32 (not AVR/Arduino Uno)",
        [("Built-in radio (WiFi) — mandatory for the wireless video stream", 0),
         ("Dual-core 240 MHz — motion on one core, comms on the other", 0),
         ("~520 KB RAM (+PSRAM on the CAM) — impossible with 2 KB on an ATmega328", 0),
         ("Hardware peripherals: RMT (step pulses) and LEDC (servo PWM)", 0),
         ("3.3 V logic — directly compatible with the ESP32-CAM, no level-shifter", 0),
         ("Programmed entirely from the Arduino IDE (C/C++) — same ecosystem", 1)])

    # 2. mechanics
    slide_section(p, "2", "Mechanics — XYZ gantry")
    slide_bullets(p, "Cartesian frame with screw actuators",
        [("3 orthogonal axes: X (left/right), Y (front/back), Z (up/down)", 0),
         ("Each axis = screw actuator (lead/ballscrew) + NEMA 17", 0),
         ("Screw rotation → translation of the nut/carriage (linear motion)", 0),
         ("Screw lead = 4 mm per turn", 1),
         ("On the Z carriage: the gripper + ESP32-CAM (load below the Z motor torque)", 0),
         ("Typical build order: Y base → X bridge → Z vertical", 0)],
        sub="Advantage: precision and repeatability (like a 3D printer / CNC)")

    # 3. motion control
    slide_section(p, "3", "Motion control")
    slide_bullets(p, "Stepper motors + A4988 drivers",
        [("NEMA 17, 1.8°/step → 200 steps per full revolution", 0),
         ("Open-loop control: count the steps ⇒ you know the position", 0),
         ("The ESP32 cannot power the motor → the A4988 driver amplifies", 0),
         ("The brain only sends: STEP (pulse = 1 step) and DIR (direction)", 1),
         ("Shared ENABLE (GPIO13, active-LOW) for all 3 drivers", 1),
         ("Runs on full step; MS1/2/3 set by jumpers, not on GPIO", 0)])

    s = blank(p); header(s, "Calibration: steps per millimeter")
    rect(s, Inches(0.8), Inches(1.6), Inches(11.7), Inches(1.4), LBLUE)
    txt(s, Inches(1.1), Inches(1.75), Inches(11.1), Inches(1.2),
        [[("steps_per_mm = (steps/turn × microstepping) / screw_lead", 22, NAVY, True)],
         [("= (200 × 1) / 4 mm  =  ", 22, DARK, False),
          ("50 steps/mm", 24, GREEN, True)]], line_spacing=1.1)
    bullets(s, Inches(0.8), Inches(3.3), Inches(11.7), Inches(3.4),
        [("Verified empirically: command 100 mm → measure with calipers → 100 mm ⇒ steps_per_mm = 50.0", 0),
         ("FastAccelStepper: timing on the RMT hardware → no lost steps at high speed", 0),
         ("3 motors at once, smooth acceleration/deceleration, without loading the CPU", 0),
         ("Motion profile: V (mm/s) and A (mm/s²) adjustable from commands", 0)])

    # 4. coordinates & homing
    slide_section(p, "4", "Coordinates & homing")
    slide_bullets(p, "Homing — establishing the origin",
        [("At power-on the physical position is unknown → need a reference zero", 0),
         ("6 endstops (MIN + MAX per axis), wired NO + INPUT_PULLUP", 0),
         ("Pressed = LOW (NC would be fail-safe, but the NC terminal did not work)", 1),
         ("Per-axis sequence: fast approach → back-off ~4 mm → slow re-approach → 0", 0),
         ("Order Z → X → Y (anti-collision: raise Z first)", 0),
         ("Z homes UP (toward MAX): the shelves are in front", 0, AMBER),
         ("Consequence: Z=0 is at the top (safe); going down toward shelves = negative", 1)])

    s = blank(p); header(s, "Software limits (soft limits)")
    add_table(s, Inches(2.0), Inches(1.9), Inches(9.3),
        [["Axis", "Min (mm)", "Max (mm)", "Note"],
         ["X", "0", "194", "homes at MIN, positive travel"],
         ["Y", "0", "170", "homes at MIN, positive travel"],
         ["Z", "−117", "0", "homes at MAX (top); down = negative"]],
        col_w=[Inches(1.3), Inches(2.2), Inches(2.2), Inches(3.6)], fsize=16)
    callout(s, Inches(2.0), Inches(4.6), Inches(9.3), Inches(1.6),
            "Safe travel",
            "Travel height is Z = −5 mm (not 0): exactly at Z=0 the MAX switch is pressed. "
            "Before any X/Y move, Z is raised to −5 mm.",
            color=RGBColor(0xFD,0xF3,0xE0), headcol=AMBER)

    # 5. gripper
    slide_section(p, "5", "Gripper")
    slide_bullets(p, "The hand that grabs the box (SG90)",
        [("Parallel gripper with 2 jaws, driven by an SG90 servo", 0),
         ("Servo = you command an angle (0–180°), it goes exactly there", 0),
         ("50 Hz PWM signal on GPIO19, generated by the ESP32 LEDC peripheral", 0),
         ("Open = 3°, closed = 47° (at 50° it strains and buzzes)", 1),
         ("Needs solid dedicated 5V: at 4.5V it only “ticked”, looked broken", 0, AMBER),
         ("Code robust to ESP32 core 2.x vs 3.x (ledcWrite/ledcAttach differ)", 1)])

    # 6. vision
    slide_section(p, "6", "Vision — reading the QR")
    slide_bullets(p, "PC-decode pipeline",
        [("The OV2640 camera captures the image (SVGA 800×600)", 0),
         ("ESP32-CAM serves an MJPEG stream over WiFi at http://<ip>/stream", 0),
         ("The PC opens the stream (OpenCV) and decodes the QR per frame", 0),
         ("pyzbar first, then OpenCV QRCodeDetector as fallback; errors caught", 1),
         ("Fixed focus (~10–20 cm): the camera is always brought to SCAN_POS", 0),
         ("On a hit, the PC sends “QR:<code>” to the brain over USB serial", 0)])

    s = blank(p); header(s, "Why decode on the PC (not on the camera)")
    callout(s, Inches(0.7), Inches(1.5), Inches(5.8), Inches(2.4),
            "On-device problem (abandoned)",
            "On-board decoding (quirc on the ESP32-CAM) gave a read rate often under 10%: "
            "fixed-focus OV2640 + a weak chip for image processing.",
            color=RGBColor(0xFD,0xE9,0xE9), headcol=RGBColor(0xB0,0x30,0x30))
    callout(s, Inches(6.8), Inches(1.5), Inches(5.8), Inches(2.4),
            "PC-decode solution",
            "The camera only streams MJPEG; the PC decodes with OpenCV/pyzbar (far more "
            "accurate) and forwards the code to the brain. Plus a live debug window.",
            color=RGBColor(0xEA,0xF7,0xEE), headcol=GREEN)
    bullets(s, Inches(0.7), Inches(4.2), Inches(12), Inches(2.4),
        [("esp32cam_stream.ino — MJPEG server on the camera (WiFi)", 0),
         ("pc/decode_qr.py — reads the stream, decodes, sends “QR:<code>” + forwards commands", 0),
         ("brainPC.ino — routes lines starting with “QR:” into the sort logic", 0),
         ("The earlier ESP-NOW on-device variant stays as a standalone fallback", 0, AMBER)])

    # 7. communication
    slide_section(p, "7", "Communication")
    s = blank(p); header(s, "Two links: camera → PC → brain")
    callout(s, Inches(0.7), Inches(1.5), Inches(5.8), Inches(2.3),
            "Link 1 — WiFi (MJPEG video)",
            "ESP32-CAM joins the WiFi, runs a small web server, serves a continuous "
            "JPEG stream at /stream. The PC reads it with cv2.VideoCapture(url).")
    callout(s, Inches(6.8), Inches(1.5), Inches(5.8), Inches(2.3),
            "Link 2 — USB serial",
            "The PC sends “QR:<code>” to the brain; the same cable carries the brain's "
            "logs back and the commands you type (HA, R1, P...).",
            color=RGBColor(0xEA,0xF7,0xEE), headcol=GREEN)
    bullets(s, Inches(0.7), Inches(4.1), Inches(12), Inches(2.4),
        [("brainPC.ino: a line starting “QR:” → sort logic; otherwise → console command", 0),
         ("decode_qr.py auto-reconnects if the stream drops", 0),
         ("Anti-spam cooldown so the same code is not re-sent too fast", 0),
         ("The camera (streamer) needs only power — no data wire to the moving carriage", 0, GREEN)])

    # 8. sorting + state machine
    slide_section(p, "8", "Sorting & state machine")
    s = blank(p); header(s, "3×2 shelves and the fixed code → cell mapping")
    add_table(s, Inches(0.7), Inches(1.7), Inches(6.0),
        [["", "Column 1", "Column 2"],
         ["Row 1", "MS0011", "MS0012"],
         ["Row 2", "CJ0011", "CJ0012"],
         ["Row 3", "EB0011", "Boxes (pickup)"]],
        col_w=[Inches(1.4), Inches(2.3), Inches(2.3)], fsize=15)
    bullets(s, Inches(7.0), Inches(1.7), Inches(5.8), Inches(4.5),
        [("CODE_MAP: fixed table code → (row, column)", 0),
         ("Strategy A (fixed mapping), not “first free”", 1),
         ("Small, fixed set ⇒ deterministic, easy to debug", 1),
         ("Each cell: (x,y,z) taught manually (teach)", 0),
         ("Positions hardcoded in the sketch (no NVS)", 1),
         ("Unknown code → rejected (no move)", 0, AMBER)])

    s = blank(p); header(s, "State machine (the AUTO cycle, R1)")
    flow = ["HOMING\nZ→X→Y", "GO TO\nSCAN_POS", "WAIT\nfor QR",
            "VALIDATE\nknown code?", "PICK\nfrom Boxes", "PLACE\nat cell", "RETURN\nto SCAN"]
    bx = Inches(0.45); bw = Inches(1.72); gap = Inches(0.08); y = Inches(2.6)
    for i, t in enumerate(flow):
        x = bx + i*(bw+gap)
        col = AMBER if "VALIDATE" in t else BLUE
        rect(s, x, y, bw, Inches(1.5), LBLUE)
        rect(s, x, y, bw, Inches(0.06), col)
        txt(s, x+Inches(0.05), y, bw-Inches(0.1), Inches(1.5),
            [[(t, 13, DARK, False)]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    bullets(s, Inches(0.7), Inches(4.5), Inches(12), Inches(2.4),
        [("Z always raised before any X/Y move (anti-collision)", 0),
         ("Unknown/missing code → rejected, the cycle does not get stuck", 0),
         ("5 s cooldown after each sort so repeated detections do not loop the cycle", 0),
         ("STOP (S or R0) instantly halts motion and auto mode", 0)])

    # 9. power + pins
    slide_section(p, "9", "Power & pins")
    s = blank(p); header(s, "Power distribution")
    bullets(s, Inches(0.8), Inches(1.7), Inches(11.7), Inches(2.8),
        [("12V → the 3 A4988 drivers (VMOT) + 100µF capacitor/driver", 0),
         ("Buck 12V→3.3V → ESP32 brain + driver logic", 0),
         ("Dedicated 5V → SG90 servo (+ cap, anti voltage-sag)", 0),
         ("ESP32-CAM: powered separately over USB 5V (laptop/power bank)", 0),
         ("COMMON GND mandatory between supply, bucks, brain, drivers, servo", 0, AMBER),
         ("The camera (WiFi) does NOT need a common GND with the brain", 1)])
    callout(s, Inches(0.8), Inches(5.0), Inches(11.7), Inches(1.3),
            "Separate domains",
            "12V power (motors) and 3.3V logic are different domains, tied only by the common GND. "
            "Capacitors absorb current spikes and prevent brownout/reset.")

    s = blank(p); header(s, "Pin map (brain ESP32 WROOM-32D)")
    add_table(s, Inches(1.3), Inches(1.6), Inches(10.7),
        [["Function", "GPIO", "Notes"],
         ["X STEP / DIR", "25 / 26", "pulses + direction, motor X"],
         ["Y STEP / DIR", "32 / 33", "motor Y"],
         ["Z STEP / DIR", "27 / 14", "motor Z"],
         ["ENABLE (all 3)", "13", "active-LOW, shared"],
         ["Endstop MIN X/Y/Z", "21 / 22 / 23", "INPUT_PULLUP, pressed=LOW"],
         ["Endstop MAX X/Y/Z", "4 / 18 / 17", "INPUT_PULLUP"],
         ["SG90 servo", "19", "PWM LEDC 50 Hz"],
         ["QR input", "USB serial", "“QR:<code>” from the PC decoder"]],
        col_w=[Inches(3.6), Inches(2.4), Inches(4.7)], fsize=13)

    # 10. decisions
    slide_section(p, "10", "Decisions & lessons")
    slide_bullets(p, "Design decisions (and WHY)",
        [("Two MCUs: isolate timing-critical motion from image handling", 0),
         ("ESP32 vs Arduino: radio + RAM + dual-core + hardware timing", 0),
         ("Decode on the PC: much higher QR read accuracy than on-device", 0),
         ("ESP-NOW on-device variant abandoned (read rate was the bottleneck)", 0),
         ("FastAccelStepper: no lost steps at speed/microstepping", 0),
         ("Fixed mapping (A): deterministic on a small code set", 0),
         ("Z homes up: avoids collision with the shelves in front", 0),
         ("Hardcoded positions: simple “teach & paste” flow for a prototype", 0)], size=18)

    s = blank(p); header(s, "Problems encountered & solutions")
    add_table(s, Inches(0.6), Inches(1.6), Inches(12.1),
        [["Problem", "Cause", "Solution"],
         ["Servo only ticked", "4.5V supply", "solid dedicated 5V + common GND"],
         ["Camera kept resetting", "weak USB / auto-off", "good USB, 470µF cap"],
         ["Low on-device QR rate", "fixed focus / weak chip", "moved decoding to the PC"],
         ["Stream read fails", "WiFi hiccup", "decode_qr.py auto-reconnects"],
         ["Cycle looped on one box", "repeated detections", "5 s cooldown after each sort"],
         ["“endstop hit” on Z", "Z=0 sits on the switch", "travel at Z=−5 mm"]],
        col_w=[Inches(3.3), Inches(3.3), Inches(5.5)], fsize=13)

    slide_bullets(p, "Conclusions",
        [("XYZ gantry + steppers/A4988, driven by an ESP32 (state machine)", 0, BLUE),
         ("ESP32-CAM as a wireless camera; the PC is the smart decoder", 0),
         ("Deterministic sorting via a fixed code → cell table", 0),
         ("Robust architecture: vision/motion split + PC-side accuracy", 0),
         ("Many decisions validated hands-on at the bench (power, decoding, Z homing)", 0, GREEN)],
        sub="Thank you! Questions?")

    save(p, "Presentation-Technical.pptx")

deck_simple()
deck_technical()
print("Done.")
